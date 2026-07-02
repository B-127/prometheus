#!/usr/bin/env python3
"""Validated, secure append of a data file into the CCPI workbook.

The master workbook never leaves your machine. An input file is validated against
the exact structure of the target tab; only if it passes is it appended, after a
timestamped backup, via an atomic write guarded by an exclusive lock.

Usage:
  python tools/append_data.py --tab data      --input new_movements.xlsx
  python tools/append_data.py --tab subgroup  --input new_subgroup.csv
  python tools/append_data.py --tab data      --input f.xlsx --dry-run
"""
from __future__ import annotations
import os
import sys
import argparse
import shutil
import datetime as dt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from schema import TAB1, TAB2, SPEC
from validate import validate, find_duplicates, ValidationError

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = os.path.join(REPO, "docs", "data", "ccpi_combined_single_tab.xlsx")
BACKUP_DIR = os.path.join(REPO, "backups")
TAB_ALIAS = {"data": TAB1, "subgroup": TAB2}
INT_COLS = {"Year", "Month Number", "Month No"}


class _Lock:
    """Cross-platform best-effort exclusive lock via O_CREAT|O_EXCL."""
    def __init__(self, target): self.lock = target + ".lock"; self.fd = None
    def __enter__(self):
        try:
            self.fd = os.open(self.lock, os.O_CREAT | os.O_EXCL | os.O_RDWR)
        except FileExistsError:
            raise ValidationError("workbook is locked by another append in progress.")
        return self
    def __exit__(self, *a):
        if self.fd is not None:
            os.close(self.fd)
        if os.path.exists(self.lock):
            os.remove(self.lock)


def _typed(col, val):
    if val is None or (isinstance(val, str) and val.strip() == ""):
        return None
    if col in INT_COLS:
        try: return int(float(val))
        except (TypeError, ValueError): return None
    if col == "Date":
        try: return pd.to_datetime(val).date()
        except (TypeError, ValueError): return None
    return val


def _coerce_date(rec):
    if "Date" in rec and (rec.get("Date") in (None, "")):
        try:
            rec["Date"] = dt.date(int(rec["Year"]), int(rec["Month Number"]), 1)
        except (TypeError, ValueError, KeyError):
            pass
    return rec


def append(tab_alias: str, input_path: str, dry_run: bool):
    if tab_alias not in TAB_ALIAS:
        raise ValidationError("--tab must be 'data' or 'subgroup'.")
    tab = TAB_ALIAS[tab_alias]
    cols = SPEC[tab]["cols"]

    clean, report = validate(input_path, tab)
    existing = pd.read_excel(MASTER, sheet_name=tab, dtype=str)
    dups = find_duplicates(clean, existing, tab)

    print(f"\n=== Validation report ({tab}) ===")
    print(f"  input rows        : {report['rows_in']}")
    print(f"  valid rows        : {report['rows_valid']}")
    print(f"  duplicate keys    : {len(dups)}" + (f"  -> {dups[:3]}..." if dups else ""))
    if dups:
        print("  NOTE: duplicates will be appended as-is (you confirmed values are pre-checked).")
    if dry_run:
        print("  DRY RUN — nothing written.\n")
        return

    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    with _Lock(MASTER):
        backup = os.path.join(BACKUP_DIR, f"ccpi_backup_{stamp}.xlsx")
        shutil.copy2(MASTER, backup)

        wb = load_workbook(MASTER)        # preserves both sheets + formatting
        ws = wb[tab]
        for _, r in clean.iterrows():
            rec = _coerce_date({c: r[c] for c in cols})
            ws.append([_typed(c, rec[c]) for c in cols])
            for cell in ws[ws.max_row]:    # consistent font on new rows only
                cell.font = Font(name="Arial")

        tmp = MASTER + ".tmp"
        wb.save(tmp)
        os.replace(tmp, MASTER)            # atomic swap

    print(f"  appended {len(clean)} row(s) to '{tab}'.")
    print(f"  backup written    : {os.path.relpath(backup, REPO)}")
    print("  commit the updated workbook to publish it to the dashboard.\n")


def main():
    ap = argparse.ArgumentParser(description="Append validated data into the CCPI workbook.")
    ap.add_argument("--tab", required=True, choices=["data", "subgroup"])
    ap.add_argument("--input", required=True)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    try:
        append(args.tab, args.input, args.dry_run)
    except ValidationError as e:
        print(f"\nREJECTED: {e}\n", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
